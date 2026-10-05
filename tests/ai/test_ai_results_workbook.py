from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]


def test_ai_workbook_contains_predictions_and_comparison():
    path = ROOT / "data" / "processed" / "ai_results.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert {"predictions_test", "comparaison_modeles"}.issubset(workbook.sheetnames)
    columns = next(workbook["predictions_test"].iter_rows(values_only=True))
    assert "y" in columns
    assert any(str(column).startswith("pred_") for column in columns)
    workbook.close()